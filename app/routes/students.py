import csv
import io
from datetime import datetime, date, timedelta
from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file, abort, Response
from flask_login import login_required, current_user
from sqlalchemy import or_
from openpyxl import load_workbook

from app.models import db, Student, Batch, BatchStudent, AttendanceRecord, AttendanceSession, Setting
from app.forms import StudentForm, StudentImportForm
from app.utils.decorators import admin_required, teacher_or_admin_required
from app.utils.audit import log_action

students_bp = Blueprint('students', __name__)


@students_bp.route('/')
@login_required
def index():
    search = request.args.get('search', '').strip()
    batch_filter = request.args.get('batch_id', type=int)
    status_filter = request.args.get('status', 'Active').strip()
    page = request.args.get('page', 1, type=int)

    query = Student.query
    if current_user.is_teacher:
        # Limit to batches assigned to this teacher
        teacher_batches = Batch.query.filter_by(teacher_id=current_user.id).all()
        t_batch_ids = [b.id for b in teacher_batches]
        query = query.join(BatchStudent).filter(BatchStudent.batch_id.in_(t_batch_ids))

    if batch_filter:
        query = query.join(BatchStudent).filter(BatchStudent.batch_id == batch_filter, BatchStudent.status == 'Active')

    if status_filter:
        query = query.filter(Student.status == status_filter)

    if search:
        query = query.filter(
            or_(
                Student.student_name.ilike(f'%{search}%'),
                Student.roll_number.ilike(f'%{search}%'),
                Student.email.ilike(f'%{search}%'),
                Student.mobile_number.ilike(f'%{search}%')
            )
        )

    pagination = query.order_by(Student.roll_number).paginate(page=page, per_page=15, error_out=False)
    
    batches_list = Batch.query.filter_by(status='Active').order_by(Batch.batch_name).all()
    if current_user.is_teacher:
        batches_list = [b for b in batches_list if b.teacher_id == current_user.id]

    threshold = float(Setting.get_value('low_attendance_threshold', 75.0))

    # Calculate quick attendance stat for students on current page
    student_stats_map = {}
    for st in pagination.items:
        student_stats_map[st.id] = st.calculate_attendance_stats()

    return render_template(
        'students/index.html',
        pagination=pagination,
        students=pagination.items,
        batches=batches_list,
        search=search,
        batch_filter=batch_filter,
        status_filter=status_filter,
        student_stats_map=student_stats_map,
        threshold=threshold
    )


@students_bp.route('/new', methods=['GET', 'POST'])
@admin_required
def create():
    form = StudentForm()
    batches = Batch.query.filter_by(status='Active').order_by(Batch.batch_name).all()
    form.batch_id.choices = [(0, '-- Do Not Assign to Batch --')] + [(b.id, f"{b.batch_code} - {b.batch_name}") for b in batches]

    if form.validate_on_submit():
        student = Student(
            roll_number=form.roll_number.data.strip().upper(),
            student_name=form.student_name.data.strip(),
            father_name=form.father_name.data.strip() if form.father_name.data else None,
            mother_name=form.mother_name.data.strip() if form.mother_name.data else None,
            mobile_number=form.mobile_number.data.strip() if form.mobile_number.data else None,
            email=form.email.data.strip().lower() if form.email.data else None,
            date_of_birth=form.date_of_birth.data,
            admission_date=form.admission_date.data,
            status=form.status.data
        )
        db.session.add(student)
        db.session.flush()

        # Batch assignment if chosen
        if form.batch_id.data and form.batch_id.data != 0:
            bs = BatchStudent(
                batch_id=form.batch_id.data,
                student_id=student.id,
                roll_number_in_batch=student.roll_number,
                joined_date=date.today(),
                status='Active'
            )
            db.session.add(bs)

        db.session.commit()
        log_action('CREATE_STUDENT', 'STUDENTS', student.id, None, f"Created student {student.roll_number}: {student.student_name}")
        flash(f'Student "{student.student_name}" created successfully.', 'success')
        return redirect(url_for('students.profile', student_id=student.id))

    return render_template('students/form.html', form=form, title='Add New Student')


@students_bp.route('/<int:student_id>')
@login_required
def profile(student_id):
    student = Student.query.get_or_404(student_id)
    batches = student.get_batches()

    # Teacher check: can only view if student belongs to their batch
    if current_user.is_teacher:
        teacher_batches = Batch.query.filter_by(teacher_id=current_user.id).all()
        teacher_batch_ids = {b.id for b in teacher_batches}
        student_batch_ids = {b.id for b in batches}
        if not (teacher_batch_ids & student_batch_ids):
            flash('Access denied: You do not have permission to view this student profile.', 'danger')
            abort(403)

    # Calculate overall stats
    stats = student.calculate_attendance_stats()
    threshold = float(Setting.get_value('low_attendance_threshold', 75.0))

    # Recent attendance history (last 20 sessions)
    recent_records = AttendanceRecord.query.filter_by(student_id=student.id)\
        .join(AttendanceSession)\
        .order_by(AttendanceSession.session_date.desc(), AttendanceSession.id.desc())\
        .limit(20).all()

    # Monthly breakdown for the current year
    current_year = date.today().year
    monthly_data = []
    for m in range(1, 13):
        start_d = date(current_year, m, 1)
        if m == 12:
            end_d = date(current_year, 12, 31)
        else:
            end_d = date(current_year, m + 1, 1) - timedelta(days=1)
            
        m_stats = student.calculate_attendance_stats(start_date=start_d, end_date=end_d)
        if m_stats['total_classes'] > 0:
            monthly_data.append({
                'month_name': start_d.strftime("%B"),
                'stats': m_stats
            })

    return render_template(
        'students/profile.html',
        student=student,
        batches=batches,
        stats=stats,
        threshold=threshold,
        recent_records=recent_records,
        monthly_data=monthly_data,
        current_year=current_year
    )


@students_bp.route('/<int:student_id>/edit', methods=['GET', 'POST'])
@admin_required
def edit(student_id):
    student = Student.query.get_or_404(student_id)
    form = StudentForm(original_student=student, obj=student)
    batches = Batch.query.filter_by(status='Active').order_by(Batch.batch_name).all()
    form.batch_id.choices = [(0, '-- No Change --')] + [(b.id, f"{b.batch_code} - {b.batch_name}") for b in batches]

    if form.validate_on_submit():
        old_val = f"{student.roll_number}: {student.student_name}"
        student.roll_number = form.roll_number.data.strip().upper()
        student.student_name = form.student_name.data.strip()
        student.father_name = form.father_name.data.strip() if form.father_name.data else None
        student.mother_name = form.mother_name.data.strip() if form.mother_name.data else None
        student.mobile_number = form.mobile_number.data.strip() if form.mobile_number.data else None
        student.email = form.email.data.strip().lower() if form.email.data else None
        student.date_of_birth = form.date_of_birth.data
        student.admission_date = form.admission_date.data
        student.status = form.status.data

        if form.batch_id.data and form.batch_id.data != 0:
            existing_enrollment = BatchStudent.query.filter_by(batch_id=form.batch_id.data, student_id=student.id).first()
            if not existing_enrollment:
                bs = BatchStudent(
                    batch_id=form.batch_id.data,
                    student_id=student.id,
                    roll_number_in_batch=student.roll_number,
                    joined_date=date.today(),
                    status='Active'
                )
                db.session.add(bs)

        db.session.commit()
        log_action('UPDATE_STUDENT', 'STUDENTS', student.id, old_val, f"{student.roll_number}: {student.student_name}")
        flash(f'Student "{student.student_name}" updated successfully.', 'success')
        return redirect(url_for('students.profile', student_id=student.id))

    return render_template('students/form.html', form=form, title=f'Edit Student: {student.student_name}', student=student)


@students_bp.route('/<int:student_id>/toggle-status', methods=['POST'])
@admin_required
def toggle_status(student_id):
    student = Student.query.get_or_404(student_id)
    student.status = 'Inactive' if student.status == 'Active' else 'Active'
    db.session.commit()
    log_action('TOGGLE_STUDENT_STATUS', 'STUDENTS', student.id, None, f"Student {student.roll_number} changed to {student.status}")
    flash(f'Student "{student.student_name}" is now {student.status}.', 'info')
    return redirect(url_for('students.index'))


@students_bp.route('/import', methods=['GET', 'POST'])
@admin_required
def bulk_import():
    form = StudentImportForm()
    batches = Batch.query.filter_by(status='Active').order_by(Batch.batch_name).all()
    form.batch_id.choices = [(0, '-- None (Do Not Enroll in Batch) --')] + [(b.id, f"{b.batch_code} - {b.batch_name}") for b in batches]

    if form.validate_on_submit():
        uploaded_file = form.file.data
        filename = uploaded_file.filename.lower()
        selected_batch_id = form.batch_id.data if form.batch_id.data != 0 else None

        success_count = 0
        skipped_count = 0
        errors = []

        try:
            if filename.endswith('.csv'):
                content = uploaded_file.read().decode('utf-8-sig', errors='replace')
                reader = csv.DictReader(io.StringIO(content))
                rows = list(reader)
            else:
                wb = load_workbook(uploaded_file, data_only=True)
                ws = wb.active
                headers = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
                rows = []
                for row_idx in range(2, ws.max_row + 1):
                    row_vals = [c.value for c in ws[row_idx]]
                    if any(row_vals):
                        row_dict = {headers[i]: (str(row_vals[i]).strip() if row_vals[i] is not None else '') for i in range(len(headers))}
                        rows.append(row_dict)

            for idx, r in enumerate(rows, start=2):
                roll = r.get('roll_number', '').strip().upper()
                name = r.get('student_name', '').strip()
                if not roll or not name:
                    errors.append(f"Row {idx}: Missing roll_number or student_name.")
                    skipped_count += 1
                    continue

                # Check existing roll number
                existing = Student.query.filter_by(roll_number=roll).first()
                if existing:
                    skipped_count += 1
                    errors.append(f"Row {idx}: Roll number '{roll}' already exists (Skipped).")
                    continue

                # Parse dates
                dob = None
                if r.get('date_of_birth'):
                    try:
                        dob = datetime.strptime(r.get('date_of_birth').strip(), '%Y-%m-%d').date()
                    except ValueError:
                        try:
                            dob = datetime.strptime(r.get('date_of_birth').strip(), '%d-%m-%Y').date()
                        except ValueError:
                            pass

                admission_d = date.today()
                if r.get('admission_date'):
                    try:
                        admission_d = datetime.strptime(r.get('admission_date').strip(), '%Y-%m-%d').date()
                    except ValueError:
                        try:
                            admission_d = datetime.strptime(r.get('admission_date').strip(), '%d-%m-%Y').date()
                        except ValueError:
                            pass

                student = Student(
                    roll_number=roll,
                    student_name=name,
                    father_name=r.get('father_name', '').strip() or None,
                    mother_name=r.get('mother_name', '').strip() or None,
                    mobile_number=r.get('mobile_number', '').strip() or None,
                    email=r.get('email', '').strip().lower() or None,
                    date_of_birth=dob,
                    admission_date=admission_d,
                    status='Active'
                )
                db.session.add(student)
                db.session.flush()

                if selected_batch_id:
                    bs = BatchStudent(
                        batch_id=selected_batch_id,
                        student_id=student.id,
                        roll_number_in_batch=student.roll_number,
                        joined_date=date.today(),
                        status='Active'
                    )
                    db.session.add(bs)

                success_count += 1

            db.session.commit()
            log_action('IMPORT_STUDENTS', 'STUDENTS', None, None, f"Imported {success_count} students. Skipped {skipped_count}.")
            flash(f'Import completed! Successfully added {success_count} students. Skipped {skipped_count}.', 'success')
            return render_template('students/import_result.html', success_count=success_count, skipped_count=skipped_count, errors=errors)

        except Exception as e:
            db.session.rollback()
            flash(f'Error processing file: {str(e)}', 'danger')

    return render_template('students/import.html', form=form)


@students_bp.route('/template')
@admin_required
def download_template():
    """Download clean CSV sample template for student bulk import."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['roll_number', 'student_name', 'father_name', 'mother_name', 'mobile_number', 'email', 'date_of_birth', 'admission_date'])
    writer.writerow(['COM-001', 'Rahul Sharma', 'Vijay Sharma', 'Sunita Sharma', '9876543210', 'rahul@example.com', '2005-04-12', '2026-08-01'])
    writer.writerow(['COM-002', 'Aman Gupta', 'Rajesh Gupta', 'Meena Gupta', '9876543211', 'aman@example.com', '2005-06-25', '2026-08-01'])
    writer.writerow(['COM-003', 'Priya Patel', 'Suresh Patel', 'Geeta Patel', '9876543212', 'priya@example.com', '2005-09-18', '2026-08-01'])
    output.seek(0)
    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={"Content-Disposition": "attachment;filename=student_import_template.csv"}
    )
