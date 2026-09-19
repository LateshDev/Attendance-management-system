import io
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


# ==========================================
# CSV EXPORTERS
# ==========================================

def generate_csv_response(filename, headers, rows):
    """Generate a downloadable CSV in-memory stream."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    output.seek(0)
    return output.getvalue().encode('utf-8-sig')


# ==========================================
# EXCEL EXPORTERS (openpyxl)
# ==========================================

def generate_excel_response(filename, sheet_title, title, metadata_list, headers, rows):
    """
    Generate a beautifully styled Excel workbook.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Styles
    title_font = Font(name='Calibri', size=16, bold=True, color='1E3A8A')
    meta_font = Font(name='Calibri', size=11, italic=True, color='4B5563')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    data_font = Font(name='Calibri', size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    # Title
    ws.append([title])
    ws.cell(row=1, column=1).font = title_font
    ws.row_dimensions[1].height = 25
    
    # Metadata
    current_row = 2
    for label, val in metadata_list:
        ws.append([f"{label}: {val}"])
        ws.cell(row=current_row, column=1).font = meta_font
        current_row += 1
    
    ws.append([])  # Blank row
    current_row += 1

    # Headers
    ws.append(headers)
    header_row_idx = current_row
    ws.row_dimensions[header_row_idx].height = 24
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    
    # Rows
    for row_data in rows:
        ws.append(row_data)
        current_row += 1
        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            # Center numeric and percentage columns
            val_str = str(cell.value or '')
            if val_str.endswith('%') or val_str.isdigit() or val_str in ['Present', 'Absent', 'Leave', 'P', 'A', 'L']:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ==========================================
# PDF EXPORTERS (ReportLab)
# ==========================================

def generate_pdf_report(title, metadata_dict, headers, rows, is_landscape=False, summary_stats=None):
    """
    Generate a clean, high-contrast, production-ready PDF report using ReportLab.
    """
    buffer = io.BytesIO()
    pagesize = landscape(letter) if is_landscape else letter
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=6
    )
    meta_style = ParagraphStyle(
        'ReportMeta',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#4B5563')
    )
    stat_style = ParagraphStyle(
        'StatSummary',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1E293B')
    )

    elements = []

    # Header Title
    elements.append(Paragraph(title, title_style))
    
    # Metadata string
    meta_text = " | ".join([f"<b>{k}:</b> {v}" for k, v in metadata_dict.items()])
    elements.append(Paragraph(meta_text, meta_style))
    elements.append(Spacer(1, 10))

    # Summary Stats Box if provided
    if summary_stats:
        stat_text = " &nbsp;&nbsp;|&nbsp;&nbsp; ".join([f"<b>{k}:</b> {v}" for k, v in summary_stats.items()])
        elements.append(Paragraph(f"<font color='#047857'>{stat_text}</font>", stat_style))
        elements.append(Spacer(1, 10))

    # Table data preparation
    table_data = [headers]
    for row in rows:
        clean_row = []
        for cell in row:
            val_str = str(cell) if cell is not None else ""
            clean_row.append(Paragraph(val_str, styles['Normal']))
        table_data.append(clean_row)

    # Style Table
    t = Table(table_data, repeatRows=1)
    t_style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F8FAFC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F1F5F9')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 4),
        ('TOPPADDING', (0, 1), (-1, -1), 4),
    ])
    t.setStyle(t_style)
    elements.append(t)

    # Footer note
    elements.append(Spacer(1, 15))
    generated_on = datetime.now().strftime("%d-%m-%Y %I:%M %p")
    footer_text = f"Report generated automatically on {generated_on} | Batch-wise Attendance Management System"
    elements.append(Paragraph(f"<i><font color='#94A3B8' size='8'>{footer_text}</font></i>", meta_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
